from __future__ import annotations

import json
import re
import sys
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter

DATAGEN_ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = DATAGEN_ROOT.parent
if str(DATAGEN_ROOT) not in sys.path:
    sys.path.insert(0, str(DATAGEN_ROOT))

from app.core.domain.conversion_rules import ALLOWED_OUTPUT_TYPES
from app.core.domain.enums import CaseMode, CharacterSet, DataType, RelationType
from app.infrastructure.converters.schema_converter import convert_to_generation_run
from app.infrastructure.errors import SchemaValidationError

DEFAULT_INPUT_PATH = PROJECT_ROOT / "datagen" / "params"
SHEET_TABLES = "tables"
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}

GENERATOR_TYPE_ALIASES = {
    "STRING": DataType.STRING.value,
    "INT": DataType.INT.value,
    "FLOAT": DataType.FLOAT.value,
    "DATE": DataType.DATE.value,
    "TIMESTAMP": DataType.TIMESTAMP.value,
    "BOOLEAN": DataType.BOOLEAN.value,
    "DATE_STRING": DataType.DATE.value,
    "TIMESTAMP_STRING": DataType.TIMESTAMP.value,
}
ALLOWED_OUTPUT_TYPES_BY_SOURCE = {
    source_type.value: {target_type.value for target_type in target_types}
    for source_type, target_types in ALLOWED_OUTPUT_TYPES.items()
}
SUPPORTED_CONSTRAINTS = {
    DataType.STRING.value: {
        "null_ratio",
        "is_unique",
        "allowed_values",
        "length",
        "regular_expr",
        "character_set",
        "case_mode",
    },
    DataType.INT.value: {
        "null_ratio",
        "is_unique",
        "allowed_values",
        "min_value",
        "max_value",
    },
    DataType.FLOAT.value: {
        "null_ratio",
        "is_unique",
        "allowed_values",
        "min_value",
        "max_value",
        "precision",
    },
    DataType.DATE.value: {
        "null_ratio",
        "is_unique",
        "allowed_values",
        "min_value",
        "max_value",
        "date_format",
    },
    DataType.TIMESTAMP.value: {
        "null_ratio",
        "is_unique",
        "allowed_values",
        "min_timestamp",
        "max_timestamp",
        "timestamp_format",
    },
    DataType.BOOLEAN.value: {
        "null_ratio",
        "is_unique",
        "allowed_values",
    },
}
HEADER_ALIASES = {
    "column_name": "column_name",
    "name": "column_name",
    "gen_data_type": "generator_data_type",
    "generator_data_type": "generator_data_type",
    "data_type": "generator_data_type",
    "output_data_type": "output_data_type",
    "is_primary_key": "is_primary_key",
    "constraints": "constraints",
    "foreign_key": "foreign_key",
}
LEGACY_CONSTRAINT_ALIASES = {
    "unique": "is_unique",
    "min_date": "min_value",
    "max_date": "max_value",
}
INT_PATTERN = re.compile(r"^[+-]?\d+$")
FLOAT_PATTERN = re.compile(r"^[+-]?(?:\d+\.\d+|\d+\.\d*|\.\d+)$")


@dataclass(frozen=True)
class ValidationIssue:
    sheet: str
    row: int | None
    column_index: int | None
    field: str | None
    message: str


@dataclass(frozen=True)
class ExcelSheetNames:
    tables: str = SHEET_TABLES


class WorkbookValidationError(ValueError):

    def __init__(self, issues: list[ValidationIssue]):
        self.issues = sorted(
            issues,
            key=lambda issue: (
                issue.sheet,
                issue.row or 0,
                issue.column_index or 0,
                issue.message,
            ),
        )
        super().__init__(f"Workbook validation failed with {len(self.issues)} issue(s)")


class CellValidationError(ValueError):

    def __init__(self, message: str, field: str | None = None, column_index: int | None = None):
        super().__init__(message)
        self.field = field
        self.column_index = column_index


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    try:
        if bool(pd.isna(value)):
            return ""
    except TypeError:
        pass
    return str(value).strip()


def is_blank(value: Any) -> bool:
    return normalize_text(value) == ""


def make_issue(
    sheet: str,
    message: str,
    row: int | None = None,
    column_index: int | None = None,
    field: str | None = None,
) -> ValidationIssue:
    return ValidationIssue(
        sheet=sheet,
        row=row,
        column_index=column_index,
        field=field,
        message=message,
    )


def format_validation_issue(issue: ValidationIssue) -> str:
    parts = [f"sheet={issue.sheet}"]
    if issue.row is not None:
        parts.append(f"row={issue.row}")
    if issue.column_index is not None:
        column_label = get_column_letter(issue.column_index)
        if issue.field:
            column_label = f"{column_label} ({issue.field})"
        parts.append(f"column={column_label}")
    elif issue.field:
        parts.append(f"field={issue.field}")
    parts.append(f"message={issue.message}")
    return " | ".join(parts)


def report_validation_issues(issues: list[ValidationIssue]) -> None:
    print(f"Validation failed: {len(issues)} issue(s)", file=sys.stderr)
    for issue in issues:
        print(format_validation_issue(issue), file=sys.stderr)


def prefix_issue(workbook_name: str, issue: ValidationIssue) -> ValidationIssue:
    return ValidationIssue(
        sheet=f"{workbook_name}:{issue.sheet}",
        row=issue.row,
        column_index=issue.column_index,
        field=issue.field,
        message=issue.message,
    )


def parse_bool(value: Any, field_name: str) -> bool | None:
    text = normalize_text(value)
    if not text:
        return None
    lowered = text.lower()
    if lowered in {"true", "1", "yes", "y"}:
        return True
    if lowered in {"false", "0", "no", "n"}:
        return False
    raise ValueError(f"Invalid boolean for {field_name}: {value!r}")


def normalize_generator_type(value: Any) -> str:
    text = normalize_text(value)
    if not text:
        raise ValueError("generator_data_type must not be empty")
    normalized = GENERATOR_TYPE_ALIASES.get(text.upper())
    if not normalized:
        raise ValueError(f"Unsupported generator_data_type: {value!r}")
    return normalized


def normalize_output_type(value: Any) -> str | None:
    text = normalize_text(value)
    if not text:
        return None
    normalized = GENERATOR_TYPE_ALIASES.get(text.upper())
    if not normalized:
        raise ValueError(f"Unsupported output_data_type: {value!r}")
    return normalized


def split_semicolon_pairs(raw_text: str) -> list[str]:
    parts: list[str] = []
    current: list[str] = []
    quote_char: str | None = None
    square_depth = 0
    curly_depth = 0

    for char in raw_text:
        if quote_char:
            current.append(char)
            if char == quote_char:
                quote_char = None
            continue

        if char in {'"', "'"}:
            quote_char = char
            current.append(char)
            continue

        if char == "[":
            square_depth += 1
        elif char == "]":
            square_depth = max(0, square_depth - 1)
        elif char == "{":
            curly_depth += 1
        elif char == "}":
            curly_depth = max(0, curly_depth - 1)

        if char == ";" and square_depth == 0 and curly_depth == 0:
            token = "".join(current).strip()
            if token:
                parts.append(token)
            current = []
            continue

        current.append(char)

    tail = "".join(current).strip()
    if tail:
        parts.append(tail)
    return parts


def parse_scalar(raw_value: str) -> Any:
    value = raw_value.strip()
    if value == "":
        return ""

    if value.startswith("[") or value.startswith("{"):
        try:
            return json.loads(value)
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSON value: {value!r}") from exc

    lowered = value.lower()
    if lowered == "true":
        return True
    if lowered == "false":
        return False
    if lowered == "null":
        return None
    if INT_PATTERN.fullmatch(value):
        return int(value)
    if FLOAT_PATTERN.fullmatch(value):
        return float(value)
    if (value.startswith('"') and value.endswith('"')) or (value.startswith("'") and value.endswith("'")):
        return value[1:-1]
    return value


def parse_key_value_text(raw_text: Any, field_name: str) -> dict[str, Any]:
    text = normalize_text(raw_text)
    if not text:
        return {}

    result: dict[str, Any] = {}
    for chunk in split_semicolon_pairs(text):
        if "=" not in chunk:
            raise ValueError(f"{field_name} contains malformed fragment: {chunk!r}")
        key, value = chunk.split("=", 1)
        normalized_key = key.strip()
        if not normalized_key:
            raise ValueError(f"{field_name} contains empty key")
        result[normalized_key] = parse_scalar(value)
    return result


def ensure_list(value: Any, field_name: str) -> list[Any]:
    if not isinstance(value, list):
        raise ValueError(f"{field_name} must be a JSON array")
    return value


def normalize_string_token(value: Any, field_name: str) -> str:
    text = normalize_text(value)
    if not text:
        raise ValueError(f"{field_name} must not be empty")
    return text


def normalize_constraints(raw_constraints: dict[str, Any], generator_type: str) -> dict[str, Any]:
    constraints: dict[str, Any] = {}
    legacy_flags: dict[str, Any] = {}

    for raw_key, raw_value in raw_constraints.items():
        key = LEGACY_CONSTRAINT_ALIASES.get(raw_key.strip(), raw_key.strip())

        if key in {"lowercase", "uppercase", "digits_only", "chars_only"}:
            legacy_flags[key] = raw_value
            continue

        if generator_type == DataType.TIMESTAMP.value and key == "date_format":
            key = "timestamp_format"

        if key not in SUPPORTED_CONSTRAINTS[generator_type]:
            supported = ", ".join(sorted(SUPPORTED_CONSTRAINTS[generator_type]))
            raise ValueError(
                f"Unsupported constraint {raw_key!r} for {generator_type}. Supported: {supported}"
            )

        if key == "null_ratio":
            if isinstance(raw_value, bool):
                raise ValueError("null_ratio must be integer from 0 to 100")
            constraints[key] = int(raw_value)
            continue

        if key == "is_unique":
            if not isinstance(raw_value, bool):
                raise ValueError("is_unique must be true/false")
            constraints[key] = raw_value
            continue

        if key == "allowed_values":
            constraints[key] = ensure_list(raw_value, "allowed_values")
            continue

        if key in {"length", "precision"}:
            if isinstance(raw_value, bool):
                raise ValueError(f"{key} must be integer")
            constraints[key] = int(raw_value)
            continue

        if key in {"min_value", "max_value"}:
            constraints[key] = raw_value
            continue

        if key in {"min_timestamp", "max_timestamp"}:
            constraints[key] = normalize_string_token(raw_value, key)
            continue

        if key in {"regular_expr", "date_format", "timestamp_format"}:
            constraints[key] = normalize_string_token(raw_value, key)
            continue

        if key == "character_set":
            value = normalize_string_token(raw_value, key).lower()
            allowed = {item.value for item in CharacterSet}
            if value not in allowed:
                raise ValueError(f"Unsupported character_set: {raw_value!r}")
            constraints[key] = value
            continue

        if key == "case_mode":
            value = normalize_string_token(raw_value, key).lower()
            allowed = {item.value for item in CaseMode}
            if value not in allowed:
                raise ValueError(f"Unsupported case_mode: {raw_value!r}")
            constraints[key] = value
            continue

        constraints[key] = raw_value

    if generator_type == DataType.STRING.value:
        lowercase = legacy_flags.get("lowercase")
        uppercase = legacy_flags.get("uppercase")
        digits_only = legacy_flags.get("digits_only")
        chars_only = legacy_flags.get("chars_only")

        if lowercase is True:
            constraints.setdefault("case_mode", CaseMode.LOWER.value)
        if uppercase is True:
            constraints.setdefault("case_mode", CaseMode.UPPER.value)
        if lowercase is True and uppercase is True:
            raise ValueError("lowercase=true and uppercase=true cannot be used together")
        if digits_only is True:
            constraints.setdefault("character_set", CharacterSet.DIGITS.value)
        if chars_only is True:
            constraints.setdefault("character_set", CharacterSet.LETTERS.value)
        if digits_only is True and chars_only is True:
            raise ValueError("digits_only=true and chars_only=true cannot be used together")

    if "greater_than" in raw_constraints or "less_than" in raw_constraints:
        raise ValueError(
            "greater_than/less_than are no longer supported in current generator contract"
        )

    return constraints


def normalize_foreign_key(
    raw_foreign_key: dict[str, Any],
    current_schema_name: str,
) -> dict[str, Any] | None:
    if not raw_foreign_key:
        return None

    schema_name = normalize_text(raw_foreign_key.get("schema_name")) or current_schema_name
    table_name = normalize_text(raw_foreign_key.get("table_name"))
    column_name = normalize_text(raw_foreign_key.get("column_name"))
    relation_type = normalize_text(raw_foreign_key.get("relation_type")).upper()

    if not table_name or not column_name or not relation_type:
        raise ValueError("foreign_key must contain table_name, column_name and relation_type")

    allowed_relation_types = {item.value for item in RelationType}
    if relation_type not in allowed_relation_types:
        raise ValueError(f"Unsupported relation_type: {relation_type!r}")

    return {
        "schema_name": schema_name,
        "table_name": table_name,
        "column_name": column_name,
        "relation_type": relation_type,
    }


def infer_schema_error_field(message: str) -> str | None:
    lowered = message.lower()
    if "relation_type" in lowered:
        return "foreign_key"
    if "output data type" in lowered or "unsupported conversion" in lowered:
        return "output_data_type"
    if "generator data type" in lowered:
        return "generator_data_type"
    if "foreign key column" in lowered:
        return "constraints"
    if "primary key" in lowered and "null_ratio" in lowered:
        return "constraints"
    if "primary key" in lowered:
        return "is_primary_key"
    constraint_markers = (
        "null_ratio",
        "allowed_values",
        "case_mode",
        "character_set",
        "date_format",
        "timestamp_format",
        "min_value",
        "max_value",
        "min_timestamp",
        "max_timestamp",
        "precision",
        "length",
        "regular_expr",
    )
    if any(marker in lowered for marker in constraint_markers):
        return "constraints"
    return None


class ExcelToDictConverter:

    def __init__(self, excel_path: str | Path, sheet_names: ExcelSheetNames | None = None) -> None:
        self._excel_path = Path(excel_path)
        self._sheet_names = sheet_names or ExcelSheetNames()

    def convert(self, validate_schema: bool = True) -> dict[str, Any]:
        warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")
        try:
            xls = pd.ExcelFile(self._excel_path)
        except FileNotFoundError as exc:
            raise WorkbookValidationError(
                [make_issue(sheet="workbook", message=f"Workbook file not found: {self._excel_path}")]
            ) from exc
        except ValueError as exc:
            raise WorkbookValidationError(
                [make_issue(sheet="workbook", message=f"Invalid Excel file: {self._excel_path}")]
            ) from exc

        raw_tables_df = self._read_sheet(xls, self._sheet_names.tables, required=True, header=None)
        source_name = self._find_source_name(raw_tables_df)
        table_specs = self._build_table_specs(raw_tables_df)

        issues: list[ValidationIssue] = []
        entities: list[dict[str, Any]] = []

        for table_spec in table_specs:
            sheet_name = table_spec["sheet_name"]
            df = self._read_sheet(xls, sheet_name, required=True)
            normalized_df, positions = self._normalize_data_sheet(df)
            required_headers = {"column_name", "generator_data_type", "constraints", "foreign_key", "is_primary_key"}
            missing_headers = sorted(required_headers - set(positions))
            if missing_headers:
                issues.append(
                    make_issue(
                        sheet=sheet_name,
                        row=1,
                        message=f"Sheet is missing columns: {', '.join(missing_headers)}",
                    )
                )
                continue

            columns, column_issues = self._build_columns(
                table_spec=table_spec,
                df=normalized_df,
                positions=positions,
                validate_schema=validate_schema,
            )
            issues.extend(column_issues)
            if columns:
                entities.append(
                    {
                        "schema_name": table_spec["schema_name"],
                        "table_name": table_spec["table_name"],
                        "total_rows": table_spec["total_rows"],
                        "columns": columns,
                    }
                )

        if issues:
            raise WorkbookValidationError(issues)

        payload = {
            "source_name": source_name,
            "entities": entities,
        }
        if validate_schema:
            try:
                convert_to_generation_run(run_id="excel_import_validation", raw_tables=payload["entities"])
            except (SchemaValidationError, ValueError) as exc:
                raise WorkbookValidationError([make_issue(sheet="workbook", message=str(exc))]) from exc
        return payload

    @staticmethod
    def _read_sheet(
        xls: pd.ExcelFile,
        sheet_name: str,
        required: bool,
        header: int | None = 0,
    ) -> pd.DataFrame:
        if sheet_name not in xls.sheet_names:
            raise WorkbookValidationError(
                [make_issue(sheet=sheet_name, message="Required sheet is missing")]
            )
        return pd.read_excel(xls, sheet_name=sheet_name, header=header, dtype=object)

    @staticmethod
    def _find_source_name(raw_tables_df: pd.DataFrame) -> str:
        max_rows = min(len(raw_tables_df.index), 10)
        max_cols = min(len(raw_tables_df.columns), 6)
        for row_index in range(max_rows):
            for column_index in range(max_cols):
                current = normalize_text(raw_tables_df.iat[row_index, column_index]).lower()
                if current != "source_name":
                    continue

                same_row_value = ""
                next_row_value = ""
                if column_index + 1 < len(raw_tables_df.columns):
                    same_row_value = normalize_text(raw_tables_df.iat[row_index, column_index + 1])
                if row_index + 1 < len(raw_tables_df.index):
                    next_row_value = normalize_text(raw_tables_df.iat[row_index + 1, column_index])

                if same_row_value:
                    return same_row_value
                if next_row_value:
                    return next_row_value
        return "excel_source"

    @staticmethod
    def _build_table_specs(raw_tables_df: pd.DataFrame) -> list[dict[str, Any]]:
        header_row = None
        header_map: dict[str, int] = {}
        max_rows = min(len(raw_tables_df.index), 20)
        max_cols = min(len(raw_tables_df.columns), 8)

        for row_index in range(max_rows):
            current_headers: dict[str, int] = {}
            for column_index in range(max_cols):
                value = normalize_text(raw_tables_df.iat[row_index, column_index]).lower()
                if value:
                    current_headers[value] = column_index
            if {"schema_name", "table_name", "total_rows"} <= set(current_headers):
                header_row = row_index
                header_map = current_headers
                break

        if header_row is None:
            raise WorkbookValidationError(
                [make_issue(sheet=SHEET_TABLES, message="Could not find table metadata header on 'tables' sheet")]
            )

        table_specs: list[dict[str, Any]] = []
        for row_index in range(header_row + 1, len(raw_tables_df.index)):
            schema_name = normalize_text(raw_tables_df.iat[row_index, header_map["schema_name"]])
            table_name = normalize_text(raw_tables_df.iat[row_index, header_map["table_name"]])
            total_rows_raw = raw_tables_df.iat[row_index, header_map["total_rows"]]

            if not schema_name and not table_name and is_blank(total_rows_raw):
                continue
            if not schema_name or not table_name or is_blank(total_rows_raw):
                raise WorkbookValidationError(
                    [make_issue(sheet=SHEET_TABLES, row=row_index + 1, message="Invalid table row on 'tables' sheet")]
                )

            try:
                total_rows = int(total_rows_raw)
            except (TypeError, ValueError) as exc:
                raise WorkbookValidationError(
                    [make_issue(sheet=SHEET_TABLES, row=row_index + 1, field="total_rows", message="total_rows must be integer")]
                ) from exc

            table_specs.append(
                {
                    "schema_name": schema_name,
                    "table_name": table_name,
                    "sheet_name": f"{schema_name}.{table_name}",
                    "total_rows": total_rows,
                }
            )

        if not table_specs:
            raise WorkbookValidationError(
                [make_issue(sheet=SHEET_TABLES, message="No table metadata found on 'tables' sheet")]
            )

        return table_specs

    @staticmethod
    def _normalize_data_sheet(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, int]]:
        normalized_df = df.dropna(how="all").copy()
        normalized_headers = []
        for header in normalized_df.columns:
            lowered = normalize_text(header).lower()
            normalized_headers.append(HEADER_ALIASES.get(lowered, lowered))
        normalized_df.columns = normalized_headers
        positions = {column_name: index + 1 for index, column_name in enumerate(normalized_df.columns) if column_name}
        return normalized_df, positions

    @staticmethod
    def _row_has_user_data(row: pd.Series) -> bool:
        relevant = (
            "column_name",
            "generator_data_type",
            "output_data_type",
            "is_primary_key",
            "constraints",
            "foreign_key",
        )
        return any(normalize_text(row.get(field)) for field in relevant)

    @staticmethod
    def _raise_cell_error(message: str, field: str | None, column_index: int | None) -> None:
        raise CellValidationError(message=message, field=field, column_index=column_index)

    def _build_columns(
        self,
        table_spec: dict[str, Any],
        df: pd.DataFrame,
        positions: dict[str, int],
        validate_schema: bool,
    ) -> tuple[list[dict[str, Any]], list[ValidationIssue]]:
        columns: list[dict[str, Any]] = []
        issues: list[ValidationIssue] = []
        saw_user_rows = False

        for row_index, row in df.iterrows():
            if not self._row_has_user_data(row):
                continue
            saw_user_rows = True
            excel_row = int(row_index) + 2
            try:
                column_payload = self._build_column_payload(
                    row=row,
                    positions=positions,
                    schema_name=table_spec["schema_name"],
                )
                if validate_schema:
                    self._validate_single_column_payload(
                        table_spec=table_spec,
                        positions=positions,
                        column_payload=column_payload,
                    )
            except CellValidationError as exc:
                issues.append(
                    make_issue(
                        sheet=table_spec["sheet_name"],
                        row=excel_row,
                        column_index=exc.column_index,
                        field=exc.field,
                        message=str(exc),
                    )
                )
                continue
            columns.append(column_payload)

        if not columns and not saw_user_rows:
            issues.append(make_issue(sheet=table_spec["sheet_name"], message="Sheet does not contain any columns"))

        return columns, issues

    def _build_column_payload(
        self,
        row: pd.Series,
        positions: dict[str, int],
        schema_name: str,
    ) -> dict[str, Any]:
        column_name = normalize_text(row.get("column_name"))
        if not column_name:
            self._raise_cell_error("column_name must not be empty", "column_name", positions.get("column_name"))

        try:
            generator_type = normalize_generator_type(row.get("generator_data_type"))
        except ValueError as exc:
            self._raise_cell_error(str(exc), "generator_data_type", positions.get("generator_data_type"))

        try:
            output_type = normalize_output_type(row.get("output_data_type"))
        except ValueError as exc:
            self._raise_cell_error(str(exc), "output_data_type", positions.get("output_data_type"))

        if output_type and output_type not in ALLOWED_OUTPUT_TYPES_BY_SOURCE[generator_type]:
            self._raise_cell_error(
                f"Unsupported conversion for {column_name}: {generator_type} -> {output_type}",
                "output_data_type",
                positions.get("output_data_type"),
            )

        try:
            is_primary_key = parse_bool(row.get("is_primary_key"), "is_primary_key") or False
        except ValueError as exc:
            self._raise_cell_error(str(exc), "is_primary_key", positions.get("is_primary_key"))

        try:
            raw_constraints = parse_key_value_text(row.get("constraints"), "constraints")
            constraints = normalize_constraints(raw_constraints, generator_type)
        except ValueError as exc:
            self._raise_cell_error(str(exc), "constraints", positions.get("constraints"))

        try:
            raw_foreign_key = parse_key_value_text(row.get("foreign_key"), "foreign_key")
            foreign_key = normalize_foreign_key(raw_foreign_key, schema_name)
        except ValueError as exc:
            self._raise_cell_error(str(exc), "foreign_key", positions.get("foreign_key"))

        payload: dict[str, Any] = {
            "name": column_name,
            "generator_data_type": generator_type,
            "constraints": constraints,
        }
        if output_type and output_type != generator_type:
            payload["output_data_type"] = output_type
        if is_primary_key:
            payload["is_primary_key"] = True
        if foreign_key:
            payload["foreign_key"] = foreign_key
        return payload

    def _validate_single_column_payload(
        self,
        table_spec: dict[str, Any],
        positions: dict[str, int],
        column_payload: dict[str, Any],
    ) -> None:
        raw_table = {
            "schema_name": table_spec["schema_name"],
            "table_name": table_spec["table_name"],
            "total_rows": table_spec["total_rows"],
            "columns": [column_payload],
        }
        try:
            convert_to_generation_run(run_id="excel_import_validation", raw_tables=[raw_table])
        except (SchemaValidationError, ValueError) as exc:
            field = infer_schema_error_field(str(exc))
            self._raise_cell_error(str(exc), field, positions.get(field) if field else None)


def find_excel_files(input_path: Path) -> list[Path]:
    if input_path.is_dir():
        files = sorted(
            path for path in input_path.iterdir()
            if path.is_file() and path.suffix.lower() in EXCEL_SUFFIXES and not path.name.startswith("~$")
        )
        if not files:
            raise WorkbookValidationError(
                [make_issue(sheet="workbook", message=f"No Excel files found in directory: {input_path}")]
            )
        return files

    if input_path.suffix.lower() not in EXCEL_SUFFIXES:
        raise WorkbookValidationError(
            [make_issue(sheet="workbook", message=f"Unsupported input path: {input_path}")]
        )
    return [input_path]


def load_payloads_from_path(input_path: str | Path, validate_schema: bool = True) -> list[dict[str, Any]]:
    path = Path(input_path)
    excel_files = find_excel_files(path)
    multiple_files = len(excel_files) > 1
    payloads: list[dict[str, Any]] = []
    issues: list[ValidationIssue] = []
    seen_tables: dict[tuple[str, str], str] = {}

    for excel_path in excel_files:
        converter = ExcelToDictConverter(excel_path)
        try:
            payload = converter.convert(validate_schema=validate_schema)
        except WorkbookValidationError as exc:
            workbook_issues = [prefix_issue(excel_path.name, issue) for issue in exc.issues] if multiple_files else exc.issues
            issues.extend(workbook_issues)
            continue

        for entity in payload["entities"]:
            key = (entity["schema_name"], entity["table_name"])
            previous_workbook = seen_tables.get(key)
            if previous_workbook:
                issue = make_issue(
                    sheet=f"{entity['schema_name']}.{entity['table_name']}",
                    message=f"Duplicate table definition also found in workbook {previous_workbook}",
                )
                issues.append(prefix_issue(excel_path.name, issue) if multiple_files else issue)
                continue
            seen_tables[key] = excel_path.name

        payloads.append(payload)

    if issues:
        raise WorkbookValidationError(issues)
    return payloads


def load_raw_tables_from_path(input_path: str | Path, validate_schema: bool = True) -> list[dict[str, Any]]:
    payloads = load_payloads_from_path(input_path=input_path, validate_schema=validate_schema)
    return [entity for payload in payloads for entity in payload["entities"]]


def resolve_raw_tables(
    raw_tables: list[dict[str, Any]] | None = None,
    input_path: str | Path = DEFAULT_INPUT_PATH,
    validate_schema: bool = True,
) -> list[dict[str, Any]]:
    if raw_tables is not None:
        return raw_tables
    return load_raw_tables_from_path(input_path=input_path, validate_schema=validate_schema)
